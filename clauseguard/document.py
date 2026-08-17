from __future__ import annotations

import importlib
import re
import zipfile
from pathlib import Path

from clauseguard.schemas import LoadedDocument


class DocumentLoadError(RuntimeError):
    pass


class DocumentLoader:
    supported_extensions = {".txt", ".docx", ".pdf"}
    max_file_bytes = 25 * 1024 * 1024
    max_extracted_characters = 5_000_000
    max_docx_entries = 10_000
    max_docx_expanded_bytes = 50 * 1024 * 1024
    max_pdf_pages = 2_000

    def load(self, file_path: str | Path) -> LoadedDocument:
        path = Path(file_path)
        if not path.exists():
            raise DocumentLoadError(f"Document not found: {path}")
        size_bytes = path.stat().st_size
        if size_bytes > self.max_file_bytes:
            raise DocumentLoadError(
                f"Document '{path.name}' exceeds the {self.max_file_bytes // (1024 * 1024)} MB safety limit."
            )
        suffix = path.suffix.lower()
        if suffix not in self.supported_extensions:
            supported = ", ".join(sorted(self.supported_extensions))
            raise DocumentLoadError(f"Unsupported file type '{suffix}'. Supported: {supported}")

        if suffix == ".txt":
            text = path.read_text(encoding="utf-8", errors="replace")
        elif suffix == ".docx":
            text = self._load_docx(path)
        else:
            text = self._load_pdf(path)

        text = self._normalize_text(text)
        if not text.strip():
            raise DocumentLoadError(f"No extractable text found in: {path}")
        if len(text) > self.max_extracted_characters:
            raise DocumentLoadError(
                f"Document '{path.name}' exceeds the extracted-text safety limit."
            )

        return LoadedDocument(
            path=str(path),
            file_type=suffix.lstrip("."),
            title=self._extract_title(text, path),
            text=text,
            metadata={"size_bytes": size_bytes},
        )

    def _load_docx(self, path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentLoadError("python-docx is required to read .docx files.") from exc

        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > self.max_docx_entries:
                    raise DocumentLoadError(
                        f"Could not read DOCX document '{path.name}': container exceeds the entry safety limit."
                    )
                expanded_bytes = sum(entry.file_size for entry in entries)
                if expanded_bytes > self.max_docx_expanded_bytes:
                    raise DocumentLoadError(
                        f"Could not read DOCX document '{path.name}': expanded content exceeds the safety limit."
                    )
                if any(entry.flag_bits & 0x1 for entry in entries):
                    raise DocumentLoadError(
                        f"Could not read DOCX document '{path.name}': encrypted containers are not supported."
                    )
            document = Document(str(path))
            paragraphs = [
                paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
            ]
            return "\n\n".join(paragraphs)
        except DocumentLoadError:
            raise
        except Exception as exc:
            raise DocumentLoadError(
                f"Could not read DOCX document '{path.name}': invalid or corrupted file."
            ) from exc

    def _load_pdf(self, path: Path) -> str:
        try:
            pdf_library = importlib.import_module("pypdf")
        except ModuleNotFoundError as exc:
            try:
                pdf_library = importlib.import_module("PyPDF2")
            except ModuleNotFoundError:
                raise DocumentLoadError("pypdf is required to read .pdf files.") from exc

        try:
            pages = []
            seen_long_pages: set[str] = set()
            with path.open("rb") as handle:
                reader = pdf_library.PdfReader(handle)
                if reader.is_encrypted:
                    raise DocumentLoadError(
                        f"Could not read PDF document '{path.name}': encrypted PDFs are not supported."
                    )
                if len(reader.pages) > self.max_pdf_pages:
                    raise DocumentLoadError(
                        f"Could not read PDF document '{path.name}': page count exceeds the safety limit."
                    )
                extracted_characters = 0
                for page in reader.pages:
                    page_text = page.extract_text() or ""
                    extracted_characters += len(page_text)
                    if extracted_characters > self.max_extracted_characters:
                        raise DocumentLoadError(
                            f"Could not read PDF document '{path.name}': extracted text exceeds the safety limit."
                        )
                    normalized_page = re.sub(r"\s+", " ", page_text).strip()
                    # Some generated PDFs expose the entire document as an identical
                    # hidden text layer on every page. Retain short repeated pages but
                    # collapse exact long-page clones before clause extraction.
                    if len(normalized_page) >= 2000:
                        if normalized_page in seen_long_pages:
                            continue
                        seen_long_pages.add(normalized_page)
                    pages.append(page_text)
            return "\n\n".join(pages)
        except DocumentLoadError:
            raise
        except Exception as exc:
            raise DocumentLoadError(
                f"Could not read PDF document '{path.name}': invalid or corrupted file."
            ) from exc

    def _extract_title(self, text: str, path: Path) -> str:
        legal_title_pattern = re.compile(
            r"\b(agreement|contract|amendment|policy|terms|statement|schedule|addendum|notice)\b",
            re.IGNORECASE,
        )
        for raw_line in text.splitlines()[:25]:
            line = raw_line.strip()
            if len(line) < 5:
                continue
            if legal_title_pattern.search(line) or line.isupper():
                return line[:160]
        return path.stem

    def _normalize_text(self, text: str) -> str:
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()
